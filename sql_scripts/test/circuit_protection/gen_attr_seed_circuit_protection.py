#!/usr/bin/env python3
"""从 circuit_protection_schema.xlsx 生成 test dim_attr_schema 种子（xlsx 为真源）。

产出: seed/dim_attr_schema_circuit_protection.csv

与 prod 导出合并策略:
  - xlsx 行为主；同 key 保留旧 CSV 中含 DK/ICPDF 运维 note 及 min/max bound
  - 旧 CSV 中 xlsx 未覆盖的行（如 lead_free）保留

用法:
    python gen_attr_seed_circuit_protection.py
    python gen_attr_seed_circuit_protection.py --xlsx path/to/circuit_protection_schema.xlsx
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = Path(__file__).resolve().parent
DEFAULT_XLSX = Path(r"e:\hardware_schema\schema_result\schemas_patched\xlsx\circuit_protection_schema.xlsx")
OUT_CSV = HERE / "seed" / "dim_attr_schema_circuit_protection.csv"
CLASSIFY_CSV = HERE / "seed" / "dim_l3_classify_circuit_protection.csv"

SCHEMA_VER = "v1.16.01"
L1_CODE = "circuit_protection"

SHEET_L2_MAP: dict[str, str] = {
    "Overcurrent_Overtemperature_Pro": "overcurrent_overtemperature_protection",
    "Semiconductor_Transient_Suppres": "semiconductor_transient_suppression",
    "Passive_Surge_Diversion_Base": "passive_surge_diversion",
    "Surge_Protection_Module_Base": "surge_protection_module",
}

# xlsx 分类英文名 → ETL l3_code（对齐 dim_l3_classify note raw=）
L3_MAP: dict[str, str] = {
    "Fuse": "fuse",
    "PPTC_Resettable_Fuse": "pptc_resettable_fuse",
    "Thermal_Cutoff": "thermal_cutoff",
    "Circuit_Breaker": "circuit_breaker",
    "TVS_Diode": "tvs_diode",
    "ESD_Suppressor": "esd_suppressor",
    "TSPD": "tspd",
    "MOV": "mov",
    "GDT": "gdt",
    "SPD_Module": "spd_module",
}

MODULE_MAP: dict[str, tuple[str, str]] = {
    "交易参数(Purchasing Attributes)": ("交易参数", "purchasing_attributes"),
    "合规参数(Regulatory Attributes)": ("合规参数", "regulatory_attributes"),
    "封装参数(Package Attributes)": ("封装参数", "package_attributes"),
    "技术参数(Tech Specs)": ("技术参数", "tech_specs"),
}

VALID_DB_TYPES = {"VARCHAR", "DOUBLE", "INT", "BOOLEAN", "ENUM"}

SCHEMA_COLS = [
    "schema_version", "l1_code", "scope_level", "scope_code",
    "std_attr_code", "std_attr_cn", "unit_std", "db_type", "precision",
    "value_domain", "min_bound", "max_bound",
    "is_l2_common", "display_ord", "attr_category_cn", "attr_category_en", "note",
]


def to_snake(field_en: str) -> str:
    s = field_en.strip().replace("-", "_")
    if "_" in s and s == s.lower():
        return s.lower()
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", s)
    return s.lower()


def normalize_row(*, db_type: str, precision: str, unit: str, value_domain: str) -> tuple[str, str, str, str]:
    t = (db_type or "VARCHAR").strip()
    u = (unit or "-").strip()
    p = (precision or "-").strip()
    vd = (value_domain or "").strip()
    if t not in VALID_DB_TYPES:
        t = u if u in VALID_DB_TYPES else "VARCHAR"
        u = "-"
    if t in ("BOOLEAN", "ENUM"):
        t = "VARCHAR"
    if u in ("-", ""):
        u = ""
    if p in ("-", ""):
        p = ""
    return t, u, p, vd


def module_cats(module: str) -> tuple[str, str]:
    return MODULE_MAP.get((module or "").strip(), ("技术参数", "tech_specs"))


def l3_code_from_row(row) -> str | None:
    name_en = str(row[2] or "").strip()
    if name_en in L3_MAP:
        return L3_MAP[name_en]
    return to_snake(name_en) if name_en else None


def parse_sheet(ws, l2_code: str, seen: set[tuple[str, str, str]], rows: list[dict]) -> None:
    l2_ord = 0
    l3_ord: dict[str, int] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        level = str(row[0]).strip()
        module = str(row[3] or "")
        field_cn = str(row[4] or "").strip()
        field_en = str(row[5] or "").strip()
        note = str(row[6] or "").strip()
        db_type, unit, precision, value_domain = normalize_row(
            db_type=str(row[7] or ""),
            precision=str(row[8] or ""),
            unit=str(row[9] or ""),
            value_domain=str(row[10] or ""),
        )
        if not field_en:
            continue
        std_code = to_snake(field_en)
        cat_cn, cat_en = module_cats(module)

        if level == "L2":
            l2_ord += 1
            key = ("l2", l2_code, std_code)
            if key in seen:
                continue
            seen.add(key)
            rows.append({
                "schema_version": SCHEMA_VER,
                "l1_code": L1_CODE,
                "scope_level": "l2",
                "scope_code": l2_code,
                "std_attr_code": std_code,
                "std_attr_cn": field_cn,
                "unit_std": unit,
                "db_type": db_type,
                "precision": precision,
                "value_domain": value_domain,
                "min_bound": "",
                "max_bound": "",
                "is_l2_common": 1,
                "display_ord": l2_ord * 5,
                "attr_category_cn": cat_cn,
                "attr_category_en": cat_en,
                "note": note[:500] if note else "",
            })
        elif level == "L3":
            l3_code = l3_code_from_row(row)
            if not l3_code:
                print(f"  ⚠ 未映射 L3: {row[1]!r} / {row[2]!r}", flush=True)
                continue
            l3_ord[l3_code] = l3_ord.get(l3_code, 0) + 1
            key = ("l3", l3_code, std_code)
            if key in seen:
                continue
            seen.add(key)
            rows.append({
                "schema_version": SCHEMA_VER,
                "l1_code": L1_CODE,
                "scope_level": "l3",
                "scope_code": l3_code,
                "std_attr_code": std_code,
                "std_attr_cn": field_cn,
                "unit_std": unit,
                "db_type": db_type,
                "precision": precision,
                "value_domain": value_domain,
                "min_bound": "",
                "max_bound": "",
                "is_l2_common": 0,
                "display_ord": l3_ord[l3_code] * 5 + 500,
                "attr_category_cn": cat_cn,
                "attr_category_en": cat_en,
                "note": note[:500] if note else "",
            })


def gen_from_xlsx(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows: list[dict] = []
    seen: set[tuple[str, str, str]] = set()
    for sheet_name in wb.sheetnames:
        l2_code = SHEET_L2_MAP.get(sheet_name)
        if not l2_code:
            print(f"  ⚠ 未映射 sheet: {sheet_name}", flush=True)
            continue
        parse_sheet(wb[sheet_name], l2_code, seen, rows)
    wb.close()
    rows.sort(key=lambda r: (r["scope_level"], r["scope_code"], r["display_ord"], r["std_attr_code"]))
    return rows


def load_existing(path: Path) -> dict[tuple[str, str, str], dict]:
    if not path.exists():
        return {}
    out: dict[tuple[str, str, str], dict] = {}
    with path.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            k = (r["scope_level"], r["scope_code"], r["std_attr_code"])
            out[k] = r
    return out


def merge_with_existing(xlsx_rows: list[dict], existing: dict[tuple[str, str, str], dict]) -> list[dict]:
    merged: list[dict] = []
    seen: set[tuple[str, str, str]] = set()
    for r in xlsx_rows:
        k = (r["scope_level"], r["scope_code"], r["std_attr_code"])
        seen.add(k)
        if k in existing:
            o = existing[k]
            note = str(o.get("note") or "")
            if note and any(x in note for x in ("DK", "ICPDF", "覆盖", "迭代")):
                r["note"] = note
            for fld in ("min_bound", "max_bound", "precision"):
                if o.get(fld) not in (None, ""):
                    r[fld] = o[fld]
        merged.append(r)
    for k, o in existing.items():
        if k not in seen:
            merged.append({fld: o.get(fld, "") for fld in SCHEMA_COLS})
    merged.sort(key=lambda r: (r["scope_level"], r["scope_code"], int(r["display_ord"] or 0), r["std_attr_code"]))
    return merged


def summarize(rows: list[dict]) -> None:
    from collections import Counter

    l2 = Counter(r["scope_code"] for r in rows if r["scope_level"] == "l2")
    l3 = Counter(r["scope_code"] for r in rows if r["scope_level"] == "l3")
    print(f"  合计: {len(rows)} 行 · L2={sum(l2.values())} / L3={sum(l3.values())}", flush=True)
    print(f"  L3 节点: {len(l3)} · 字段分布: {dict(sorted(l3.items()))}", flush=True)
    if CLASSIFY_CSV.exists():
        with CLASSIFY_CSV.open(encoding="utf-8") as f:
            classify_l3 = {r["l3_code"] for r in csv.DictReader(f)}
        schema_l3 = set(l3.keys())
        missing = classify_l3 - schema_l3
        extra = schema_l3 - classify_l3
        if missing:
            print(f"  ⚠ classify 有但 schema 无: {sorted(missing)}", flush=True)
        if extra:
            print(f"  ℹ schema 有但 classify 无: {sorted(extra)}", flush=True)


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SCHEMA_COLS, lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in SCHEMA_COLS})


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    args = ap.parse_args()
    if not args.xlsx.exists():
        print(f"ERROR: xlsx not found: {args.xlsx}", flush=True)
        return 1
    print(f"读取 {args.xlsx}", flush=True)
    xlsx_rows = gen_from_xlsx(args.xlsx)
    existing = load_existing(OUT_CSV)
    merged = merge_with_existing(xlsx_rows, existing)
    write_csv(merged, OUT_CSV)
    print(f"写入 {OUT_CSV}", flush=True)
    summarize(merged)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
